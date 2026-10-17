"""Bounded in-memory XLS/XLSX reader. No macros, links or formulas are executed."""
from http.server import BaseHTTPRequestHandler
from io import BytesIO
from hashlib import sha256
from zipfile import ZipFile, BadZipFile
import json
import xlrd
import openpyxl
import defusedxml  # Ensure hardened XML parsers are available before loading XLSX.

MAX_BYTES = 4 * 1024 * 1024
MAX_ROWS = 50001
MAX_CELLS = 1200000


def parse_excel(data):
    if not data or len(data) > MAX_BYTES:
        raise ValueError('Envie um Excel de até 4 MB. Para arquivos maiores, exporte CSV.')
    sheets = []
    cells = 0
    if data.startswith(bytes.fromhex('D0CF11E0A1B11AE1')):
        book = xlrd.open_workbook(file_contents=data, on_demand=True)
        try:
            if book.nsheets > 12:
                raise ValueError('Limite de 12 abas por arquivo.')
            for name in book.sheet_names():
                sheet = book.sheet_by_name(name)
                cells += sheet.nrows * sheet.ncols
                if sheet.nrows > MAX_ROWS or sheet.ncols > 80 or cells > MAX_CELLS:
                    raise ValueError('Planilha muito grande: limite de 50.000 linhas, 80 colunas e 1,2 milhão de células.')
                rows = []
                for index in range(sheet.nrows):
                    row = []
                    for cell in sheet.row(index):
                        if cell.ctype == xlrd.XL_CELL_ERROR:
                            row.append('#ERRO_EXCEL')
                        elif cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
                            row.append(str(int(cell.value)))
                        else:
                            row.append(str(cell.value).strip() if cell.value is not None else '')
                    rows.append(row)
                if rows:
                    sheets.append({'name': name, 'rows': rows})
        finally:
            book.release_resources()
    elif data.startswith(b'PK'):
        try:
            with ZipFile(BytesIO(data)) as archive:
                members = archive.infolist()
                if len(members) > 2000 or sum(m.file_size for m in members) > 64 * 1024 * 1024 or any(m.flag_bits & 1 for m in members):
                    raise ValueError('Arquivo compactado excessivo ou protegido por senha.')
                if '[Content_Types].xml' not in archive.namelist():
                    raise ValueError('O arquivo não é uma planilha XLSX.')
            book = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False, keep_vba=False)
            try:
                if len(book.worksheets) > 12:
                    raise ValueError('Limite de 12 abas por arquivo.')
                for sheet in book.worksheets:
                    if (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > 80:
                        raise ValueError('Limite de 50.000 linhas e 80 colunas.')
                    rows = []
                    for index, cells_row in enumerate(sheet.iter_rows()):
                        cells += len(cells_row)
                        if index >= MAX_ROWS or len(cells_row) > 80 or cells > MAX_CELLS:
                            raise ValueError('Limite de linhas ou células excedido.')
                        row = []
                        for cell in cells_row:
                            value = cell.value
                            if cell.data_type in ('f', 'e'):
                                row.append('#FORMULA_OU_ERRO')
                            elif isinstance(value, (int, float)) and float(value).is_integer():
                                row.append(str(int(value)))
                            else:
                                row.append(str(value).strip() if value is not None else '')
                        rows.append(row)
                    if rows:
                        sheets.append({'name': sheet.title, 'rows': rows})
            finally:
                book.close()
        except BadZipFile as exc:
            raise ValueError('Arquivo XLSX inválido.') from exc
    else:
        raise ValueError('Use um arquivo XLS ou XLSX válido, sem senha.')
    if not sheets:
        raise ValueError('Nenhuma aba com dados encontrada.')
    return {'sheets': sheets, 'sha256': sha256(data).hexdigest()}


class handler(BaseHTTPRequestHandler):
    def log_message(self, *_args):
        pass

    def do_POST(self):
        try:
            size = int(self.headers.get('Content-Length', '0'))
            if size <= 0 or size > MAX_BYTES:
                raise ValueError('Excel deve ter até 4 MB. Use CSV para arquivos maiores.')
            payload, status = parse_excel(self.rfile.read(size)), 200
        except ValueError as exc:
            payload, status = {'error': str(exc)}, 400
        except Exception:
            payload, status = {'error': 'Não foi possível ler o Excel. Confira o arquivo ou exporte CSV.'}, 422
        body = json.dumps(payload, ensure_ascii=False).encode('utf-8')
        if len(body) > MAX_BYTES:
            status = 413
            body = json.dumps({'error': 'Dados convertidos excedem 4 MB. Exporte CSV para importar a base maior.'}).encode('utf-8')
        self.send_response(status)
        self.send_header('Content-Type', 'application/json; charset=utf-8')
        self.send_header('Cache-Control', 'no-store')
        self.send_header('X-Content-Type-Options', 'nosniff')
        self.send_header('Content-Length', str(len(body)))
        self.end_headers()
        self.wfile.write(body)
